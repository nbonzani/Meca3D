# SPDX-License-Identifier: GPL-3.0-or-later
# ruff: noqa: E501 - fichier de contenu rédactionnel (tableaux de texte)
"""Construit le guide utilisateur Meca3D (charte méca ia) : docs/guide/Meca3D_Guide_utilisateur.docx.

Usage (python-docx n'est pas une dépendance du projet : environnement éphémère) :
    uv run --with python-docx python docs/guide/construire_guide.py
Le PDF est ensuite produit par Word (champs et table des matières mis à jour) :
    uv run --with pywin32 python docs/guide/exporter_pdf.py
Les captures viennent de `uv run python tests/outils_capture_guide.py`.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ICI = Path(__file__).resolve().parent
RESSOURCES = ICI / "ressources"
CAPTURES = ICI / "captures"
CAPTURES_LOTS = ICI.parent / "captures"
SORTIE = ICI / "Meca3D_Guide_utilisateur.docx"

REFERENCE = "MIA-DOC-2026-101"
VERSION_DOC = "1.0"
VERSION_LOGICIEL = "0.1.0"
DATE = "27/10/2026"

BLUEPRINT = "1B3A6B"
CYAN = "00B4D8"
CUIVRE = "E07A2C"
PAPIER = "FAFAF7"
PROFOND = "0F2347"
GRAPHITE = "2B2B2B"
TRAME = "E4E2DC"
COTE = "8A8A8A"
LIGNE_IMPAIRE = "F5F4EF"

TITRE = "Constantia"
CORPS = "Aptos"
MONO = "Cascadia Mono"


def rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


def typographie(texte: str) -> str:
    """Espaces fines insécables avant : ; ! ? et insécables dans les guillemets français."""
    texte = re.sub(r" ([:;!?])(?=\s|$)", " \\1", texte)
    return texte.replace("« ", "« ").replace(" »", " »")


def police(run, nom: str, taille: float, couleur: str = GRAPHITE, gras: bool = False) -> None:
    run.font.name = nom
    run.font.size = Pt(taille)
    run.font.bold = gras
    run.font.color.rgb = rgb(couleur)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribut in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribut), nom)


def interlettrage(run, vingtiemes: int) -> None:
    espacement = OxmlElement("w:spacing")
    espacement.set(qn("w:val"), str(vingtiemes))
    run._element.get_or_add_rPr().append(espacement)


def ombrer(cellule, hexa: str) -> None:
    tcpr = cellule._element.get_or_add_tcPr()
    fond = OxmlElement("w:shd")
    fond.set(qn("w:val"), "clear")
    fond.set(qn("w:color"), "auto")
    fond.set(qn("w:fill"), hexa)
    tcpr.append(fond)


def bordures(tableau, hexa: str = TRAME, taille: int = 4, aucune: bool = False) -> None:
    tblpr = tableau._element.tblPr
    elements = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{cote}")
        b.set(qn("w:val"), "nil" if aucune else "single")
        b.set(qn("w:sz"), str(taille))
        b.set(qn("w:color"), hexa)
        elements.append(b)
    tblpr.append(elements)


def marges_cellules(tableau, haut_bas: int, gauche_droite: int) -> None:
    tblpr = tableau._element.tblPr
    marges = OxmlElement("w:tblCellMar")
    for cote, valeur in (
        ("top", haut_bas),
        ("bottom", haut_bas),
        ("left", gauche_droite),
        ("right", gauche_droite),
    ):
        m = OxmlElement(f"w:{cote}")
        m.set(qn("w:w"), str(valeur))
        m.set(qn("w:type"), "dxa")
        marges.append(m)
    tblpr.append(marges)


def filet(paragraphe, position: str, hexa: str = BLUEPRINT, taille: int = 4) -> None:
    ppr = paragraphe._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{position}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(taille))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), hexa)
    bdr.append(b)
    ppr.append(bdr)


def champ(paragraphe, code: str, defaut: str = "1"):
    """Champ Word (PAGE, NUMPAGES, TOC) inséré dans un paragraphe ; renvoie le run du résultat."""
    debut = paragraphe.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    debut._element.append(f1)
    instr = paragraphe.add_run()
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = f" {code} "
    instr._element.append(t)
    sep = paragraphe.add_run()
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    sep._element.append(f2)
    resultat = paragraphe.add_run(defaut)
    fin = paragraphe.add_run()
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    fin._element.append(f3)
    return [debut, instr, sep, resultat, fin]


class Guide:
    def __init__(self) -> None:
        self.doc = Document()
        self.numero_figure = 0
        self.numero_tableau = 0
        self._styles()

    # ------------------------------------------------------------------ styles et mise en page
    def _styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = CORPS
        normal.font.size = Pt(11)
        normal.font.color.rgb = rgb(GRAPHITE)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), CORPS)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for nom, fonte, taille, couleur, avant, apres in (
            ("Heading 1", TITRE, 26, BLUEPRINT, 0, 12),
            ("Heading 2", TITRE, 17, BLUEPRINT, 16, 6),
            ("Heading 3", CORPS, 11, BLUEPRINT, 12, 4),
        ):
            s = styles[nom]
            s.font.name = fonte
            s.font.size = Pt(taille)
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = rgb(couleur)
            rfonts = s.element.rPr.rFonts
            for attribut in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(attribut), fonte)
            for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rfonts.get(qn(theme)) is not None:
                    del rfonts.attrib[qn(theme)]
            s.paragraph_format.space_before = Pt(avant)
            s.paragraph_format.space_after = Pt(apres)
            s.paragraph_format.keep_with_next = True
        styles["Heading 3"].font.all_caps = True
        for nom in ("TOC 1", "TOC 2"):
            try:
                s = styles[nom]
            except KeyError:
                continue
            s.font.name = CORPS

    def _page_normale(self, section) -> None:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(25)
        section.left_margin = section.right_margin = Mm(22)
        section.header_distance = section.footer_distance = Mm(12)

    # ------------------------------------------------------------------ page de garde
    def page_de_garde(self) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        for cote in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, cote, Mm(0))
        section.header_distance = section.footer_distance = Mm(0)
        tableau = self.doc.add_table(rows=1, cols=1)
        tableau.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordures(tableau, aucune=True)
        marges_cellules(tableau, 0, 0)
        ligne = tableau.rows[0]
        ligne.height = Mm(296)
        ligne.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        cellule = ligne.cells[0]
        cellule.width = Mm(210)
        ombrer(cellule, PROFOND)
        cellule.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        def para(espace_avant: float = 0, alignement=WD_ALIGN_PARAGRAPH.CENTER):
            p = cellule.add_paragraph()
            p.alignment = alignement
            p.paragraph_format.space_before = Mm(espace_avant)
            p.paragraph_format.space_after = Pt(0)
            return p

        cellule.paragraphs[0].paragraph_format.space_after = Pt(0)
        p = para(70)
        p.add_run().add_picture(str(RESSOURCES / "meca-ia-mono-dark.png"), width=Mm(90))
        p = para(28)
        police(p.add_run("Meca3D"), TITRE, 56, PAPIER, gras=True)
        p = para(6)
        r = p.add_run("DOCUMENTATION ET GUIDE UTILISATEUR")
        police(r, MONO, 14, CYAN)
        interlettrage(r, 56)
        p = para(10)
        police(
            p.add_run(
                typographie("Mécanique des solides indéformables : de la maquette 3D aux torseurs")
            ),
            CORPS,
            13,
            PAPIER,
        )
        p = para(4)
        police(p.add_run(f"Version du logiciel {VERSION_LOGICIEL}"), CORPS, 13, PAPIER)
        p = para(82, WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Mm(22)
        for i, texte in enumerate(
            (
                f"RÉFÉRENCE  {REFERENCE}    VERSION  {VERSION_DOC}    ÉMIS LE  {DATE}",
                "DIFFUSION  PUBLIC    LOGICIEL LIBRE  GPL v3    github.com/nbonzani/Meca3D",
            )
        ):
            if i:
                p.add_run().add_break()
            police(p.add_run(texte), MONO, 10, COTE)

        nouvelle = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page_normale(nouvelle)
        nouvelle.header.is_linked_to_previous = False
        nouvelle.footer.is_linked_to_previous = False
        self._entete_pied(nouvelle)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

    def _entete_pied(self, section) -> None:
        entete = section.header
        tableau = entete.add_table(rows=1, cols=2, width=Mm(166))
        bordures(tableau, aucune=True)
        gauche, droite = tableau.rows[0].cells
        gauche.width, droite.width = Mm(83), Mm(83)
        gauche.paragraphs[0].add_run().add_picture(
            str(RESSOURCES / "meca-ia-horizontal.png"), height=Mm(12)
        )
        p = droite.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        police(
            p.add_run(f"{REFERENCE}\nMeca3D {VERSION_LOGICIEL} — Guide utilisateur"), MONO, 9, COTE
        )
        droite.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        vide = entete.paragraphs[0]
        filet(vide, "top")
        vide.paragraph_format.space_after = Pt(0)
        # le paragraphe initial de l'en-tête passe sous le tableau
        entete._element.remove(vide._element)
        entete._element.append(vide._element)

        pied = section.footer
        p = pied.paragraphs[0]
        filet(p, "top")
        p.paragraph_format.tab_stops.add_tab_stop(Mm(83), alignment=1)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(166), alignment=2)
        police(p.add_run("méca · ia"), MONO, 9, BLUEPRINT)
        police(p.add_run("\tPublic"), MONO, 8, COTE)
        police(p.add_run("\t["), MONO, 9, BLUEPRINT)
        for r in champ(p, "PAGE"):
            police(r, MONO, 9, BLUEPRINT)
        police(p.add_run("/"), MONO, 9, BLUEPRINT)
        for r in champ(p, "NUMPAGES"):
            police(r, MONO, 9, BLUEPRINT)
        police(p.add_run("]"), MONO, 9, BLUEPRINT)

    # ------------------------------------------------------------------ blocs de contenu
    def eyebrow(self, texte: str) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(texte.upper())
        police(r, MONO, 9, CYAN)
        interlettrage(r, 36)

    def h1(self, texte: str, eyebrow: str | None = None, saut: bool = True) -> None:
        if eyebrow:
            self.eyebrow(eyebrow)
            if saut:
                self.doc.paragraphs[-1].paragraph_format.page_break_before = True
        titre = self.doc.add_heading(typographie(texte), level=1)
        if saut and not eyebrow:
            titre.paragraph_format.page_break_before = True

    def h2(self, texte: str) -> None:
        self.doc.add_heading(typographie(texte), level=2)

    def h3(self, texte: str) -> None:
        self.doc.add_heading(typographie(texte), level=3)

    def accent(self) -> None:
        # un encadré ou un tableau finit par un paragraphe vide d'espacement : la barre le
        # remplace, sinon elle déborde seule sur une page
        dernier = self.doc.paragraphs[-1]
        p = dernier if not dernier.text and not dernier.runs else self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(RESSOURCES / "accent-bar.png"), width=Mm(166), height=Pt(4))

    def _riche(self, p, texte: str, taille: float = 11, couleur: str = GRAPHITE) -> None:
        """Texte avec **gras** et `code` (Cascadia Mono)."""
        for morceau in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", typographie(texte)):
            if not morceau:
                continue
            if morceau.startswith("**"):
                police(p.add_run(morceau[2:-2]), CORPS, taille, couleur, gras=True)
            elif morceau.startswith("`"):
                police(p.add_run(morceau[1:-1]), MONO, taille - 1.5, BLUEPRINT)
            else:
                police(p.add_run(morceau), CORPS, taille, couleur)

    def p(self, texte: str) -> None:
        self._riche(self.doc.add_paragraph(), texte)

    def puces(self, elements: list[str], numerotees: bool = False) -> None:
        # numérotation écrite à la main : les listes numérotées de Word se poursuivent d'une
        # liste à l'autre dans un document python-docx
        for i, e in enumerate(elements, 1):
            if numerotees:
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(7)
                p.paragraph_format.first_line_indent = Mm(-7)
                police(p.add_run(f"{i}.	"), MONO, 10, CYAN)
                p.paragraph_format.tab_stops.add_tab_stop(Mm(7))
            else:
                p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            self._riche(p, e)

    def code(self, lignes: list[str]) -> None:
        tableau = self.doc.add_table(rows=1, cols=1)
        bordures(tableau, aucune=True)
        marges_cellules(tableau, 100, 160)
        cellule = tableau.rows[0].cells[0]
        ombrer(cellule, LIGNE_IMPAIRE)
        p = cellule.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for i, ligne in enumerate(lignes):
            if i:
                p.add_run().add_break()
            police(p.add_run(ligne), MONO, 9, BLUEPRINT)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def encadre(self, titre: str, texte: str, couleur: str = CYAN) -> None:
        tableau = self.doc.add_table(rows=1, cols=1)
        bordures(tableau, aucune=True)
        marges_cellules(tableau, 100, 180)
        cellule = tableau.rows[0].cells[0]
        ombrer(cellule, LIGNE_IMPAIRE)
        tcpr = cellule._element.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        gauche = OxmlElement("w:left")
        gauche.set(qn("w:val"), "single")
        gauche.set(qn("w:sz"), "24")
        gauche.set(qn("w:color"), couleur)
        b.append(gauche)
        tcpr.append(b)
        p = cellule.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(titre.upper())
        police(r, MONO, 9, couleur if couleur != CYAN else BLUEPRINT)
        interlettrage(r, 30)
        p2 = cellule.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        self._riche(p2, texte, 10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def legende(self, texte: str) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        police(p.add_run(typographie(texte)), MONO, 9, COTE)

    def figure(self, chemin: Path, legende: str, largeur: float = 166) -> None:
        self.numero_figure += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(chemin), width=Mm(largeur))
        self.legende(f"Figure {self.numero_figure} — {legende}")

    def tableau(
        self, entetes: list[str], lignes: list[list[str]], largeurs: list[float], legende: str = ""
    ) -> None:
        t = self.doc.add_table(rows=1 + len(lignes), cols=len(entetes))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordures(t)
        marges_cellules(t, 60, 110)
        for j, texte in enumerate(entetes):
            c = t.rows[0].cells[j]
            ombrer(c, BLUEPRINT)
            pa = c.paragraphs[0]
            pa.paragraph_format.space_after = Pt(0)
            police(pa.add_run(typographie(texte)), CORPS, 10, PAPIER, gras=True)
        for i, ligne in enumerate(lignes, 1):
            for j, texte in enumerate(ligne):
                c = t.rows[i].cells[j]
                ombrer(c, PAPIER if i % 2 == 0 else LIGNE_IMPAIRE)
                pa = c.paragraphs[0]
                pa.paragraph_format.space_after = Pt(0)
                pa.paragraph_format.line_spacing = 1.05
                self._riche(pa, texte, 10)
        for ligne in t.rows:
            for j, largeur in enumerate(largeurs):
                ligne.cells[j].width = Mm(largeur)
        # en-tête répété sur chaque page
        trpr = t.rows[0]._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        trpr.append(rep)
        if legende:
            self.numero_tableau += 1
            self.legende(f"Tableau {self.numero_tableau} — {legende}")
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def table_des_matieres(self) -> None:
        self.eyebrow("Sommaire")
        self.doc.add_heading("Table des matières", level=1)
        p = self.doc.add_paragraph()
        champ(p, 'TOC \\o "1-2" \\h \\z \\u', "Mettre à jour la table des matières (F9).")

    def enregistrer(self) -> None:
        self.doc.core_properties.title = "Meca3D — Documentation et guide utilisateur"
        self.doc.core_properties.author = "Nicolas Bonzani — méca ia"
        self.doc.core_properties.subject = f"Meca3D {VERSION_LOGICIEL}"
        self.doc.save(str(SORTIE))


# =============================================================================== contenu
def contenu(g: Guide) -> None:
    g.page_de_garde()
    g.table_des_matieres()

    # ------------------------------------------------------------------------------ 1
    g.h1("Présentation", "01 — Le logiciel")
    g.p(
        "**Meca3D** est un outil pédagogique de mécanique des solides indéformables. Il fait passer "
        "l'élève d'une maquette 3D, construite dans le modeleur **SolidCraft**, à une étude "
        "mécanique complète : choix des liaisons normalisées, graphe des liaisons, mobilité et "
        "hyperstatisme, lois d'entrée, simulation cinématique ou dynamique, torseurs de liaison."
    )
    g.p(
        "La démarche reprend celle du logiciel Decade (Thomas Paviot) : **l'élève modélise, le "
        "logiciel ne devine pas les liaisons**. Meca3D propose un type de liaison à partir de la "
        "maquette, mais chaque liaison doit être validée, une par une, par l'élève."
    )
    g.h2("Chaîne de calcul")
    g.tableau(
        ["Étape", "Outil", "Rôle"],
        [
            ["Maquette", "SolidCraft 0.15.0", "Géométrie, masses, liaisons CAO, couplages"],
            ["Modélisation", "Meca3D", "Liaisons normalisées, graphe, m et h, données d'entrée"],
            ["Calcul", "Exudyn 1.13.6", "Dynamique multicorps (intégrateur generalized-α)"],
            ["Résultats", "Meca3D", "Courbes, torseurs, export CSV"],
            ["Restitution", "SolidCraft", "Animation du mouvement simulé, trajectoires"],
        ],
        [32, 40, 94],
        "Rôle de chaque outil",
    )
    g.h2("Ce que Meca3D calcule")
    g.puces(
        [
            "le **graphe des liaisons** : nombre cyclomatique μ, cycles indépendants, bâti ;",
            "la **mobilité m** et le **degré d'hyperstatisme h** par la méthode cinématique ;",
            "le **mouvement** de tous les solides : positions, vitesses, accélérations ;",
            "en dynamique, les **torseurs statiques** de toutes les liaisons et les efforts "
            "transmis par les couplages ;",
            "les grandeurs de **capteurs** : point d'un solide relativement à un autre, torseur "
            "d'une liaison.",
        ]
    )
    g.h2("Public et contexte d'usage")
    g.p(
        "Meca3D s'adresse aux élèves ingénieurs, en TP encadré comme chez eux, et aux "
        "enseignants pour la préparation des maquettes, les démonstrations et les corrigés. Il "
        "fonctionne sous Windows 11. SolidCraft est lancé à côté, maquette ouverte ; Meca3D s'y "
        "connecte, lit le mécanisme, puis renvoie le mouvement simulé dans la maquette."
    )
    g.encadre(
        "Sans SolidCraft",
        "Une étude enregistrée (`.meca3d`) contient tout le mécanisme importé. Elle s'ouvre, se "
        "modifie et se simule sans SolidCraft. Seuls l'import et l'animation exigent SolidCraft.",
    )
    # pas de barre : la page 4 est pleine, elle déborderait seule

    # ------------------------------------------------------------------------------ 2
    g.h1("Installation", "02 — Mise en route")
    g.h2("Prérequis")
    g.tableau(
        ["Élément", "Exigence"],
        [
            ["Système", "Windows 11, 64 bits"],
            ["Modeleur", "SolidCraft 0.15.0 (format `solidcraft/26`), installé séparément"],
            ["Espace disque", "environ 300 Mo pour Meca3D décompressé"],
            ["Python", "non requis pour l'exécutable"],
        ],
        [40, 126],
    )
    g.h2("Installer l'exécutable")
    g.puces(
        [
            "Télécharger `Meca3D-v0.1.0-windows-x64.zip` depuis la page des versions : "
            "github.com/nbonzani/Meca3D/releases.",
            "Décompresser l'archive dans un dossier de travail, par exemple `Documents\\Meca3D`. "
            "Éviter les dossiers synchronisés et les chemins très longs.",
            "Lancer `meca3d.exe` (interface graphique). `meca3d-cli.exe` est la version en ligne "
            "de commande (chapitre 9).",
        ],
        numerotees=True,
    )
    g.encadre(
        "Avertissement Windows SmartScreen",
        "L'exécutable n'est pas signé numériquement. À la première ouverture, Windows peut "
        "afficher « Windows a protégé votre ordinateur » : cliquer sur **Informations "
        "complémentaires**, puis **Exécuter quand même**. Un antivirus d'entreprise peut exiger "
        "l'accord du service informatique.",
        CUIVRE,
    )
    g.h2("Préparer SolidCraft")
    g.p(
        "Meca3D dialogue avec SolidCraft par son **canal de contrôle** local (TCP sur 127.0.0.1, "
        "authentifié par un jeton que SolidCraft écrit dans son dossier de contrôle). Aucun "
        "réglage n'est nécessaire : Meca3D lit le port courant dans "
        "`%LOCALAPPDATA%\\SolidCraft\\control\\current.port`, et à défaut utilise le port 27878. "
        "Pour imposer un port, lancer SolidCraft avec l'option `--control 7878`."
    )
    g.encadre(
        "Fenêtre SolidCraft visible",
        "SolidCraft ne traite qu'une requête par image affichée. Fenêtre réduite, l'animation "
        "tombe à moins de 10 images/s. Garder la fenêtre de SolidCraft visible ; Meca3D la ramène "
        "au premier plan si elle est réduite.",
    )
    g.h2("Installation depuis les sources (enseignants, développeurs)")
    g.p("Python 3.12 et l'outil uv sont nécessaires :")
    g.code(
        [
            "git clone https://github.com/nbonzani/Meca3D.git",
            "cd Meca3D",
            "uv sync",
            "uv run meca3d",
        ]
    )
    g.accent()

    # ------------------------------------------------------------------------------ 3
    g.h1("Prise en main en cinq minutes", "03 — Premier essai")
    g.p(
        "Ce premier essai n'exige pas SolidCraft. Il utilise l'exemple de la bielle-manivelle "
        "fourni dans le dossier `exemples` de Meca3D."
    )
    g.puces(
        [
            "Lancer Meca3D, puis **Fichier → Ouvrir une étude…** et choisir "
            "`exemples\\bielle_manivelle.meca3d`.",
            "Dans l'arbre, cliquer sur **Graphe** : le graphe des liaisons s'affiche avec "
            "μ = 1, m = 1, h = 0. Le mécanisme est isostatique et mobile.",
            "Cliquer sur **Simuler (dynamique)**. Le calcul tourne dans un processus séparé ; la "
            "barre d'état montre sa progression (quelques secondes).",
            "Dans l'onglet **Courbes**, choisir en ordonnée `B_piston x (m)` : la course du "
            "piston, de 0,2 à 0,4 m.",
            "Dans l'onglet **Torseurs**, déplacer le curseur temporel pour lire les efforts dans "
            "chaque liaison à chaque instant.",
            "**Fichier → Exporter les résultats en CSV…** enregistre toutes les grandeurs pour "
            "un tableur.",
        ],
        numerotees=True,
    )
    g.figure(
        CAPTURES / "g01_ouverture_graphe.png",
        "Étude de la bielle-manivelle ouverte : graphe des liaisons",
    )
    g.figure(CAPTURES / "g06_courbes.png", "Course du piston après simulation dynamique")
    g.accent()

    # ------------------------------------------------------------------------------ 4
    g.h1("L'interface", "04 — Fenêtre principale")
    g.p("La fenêtre principale suit la démarche de l'étude, de gauche à droite :")
    g.tableau(
        ["Zone", "Contenu"],
        [
            ["Barre de menus", "Fichier, SolidCraft, Simulation, Enseignant"],
            [
                "Barre d'outils",
                "Commandes courantes : ouvrir, enregistrer, connecter, importer, simuler, annuler, exporter",
            ],
            [
                "Arbre de l'étude (gauche)",
                "Trois branches : Modèle cinématique, Données d'entrée, Résultats",
            ],
            ["Onglets (centre)", "Graphe des liaisons, Courbes, Torseurs, Animation, Journal"],
            ["Propriétés (droite)", "Réglages de l'élément sélectionné dans l'arbre"],
            [
                "Barre d'état (bas)",
                "Derniers messages, progression du calcul, m, h, état de la connexion SolidCraft",
            ],
        ],
        [46, 120],
        "Zones de la fenêtre",
    )
    g.h2("L'arbre de l'étude")
    g.tableau(
        ["Branche", "Rubriques", "Rôle"],
        [
            [
                "Modèle cinématique",
                "Solides, Liaisons, Graphe",
                "Classes d'équivalence, liaisons normalisées, graphe et m, h",
            ],
            [
                "Données d'entrée",
                "Pilotes, Pesanteur, Actions, Ressorts",
                "Lois d'entrée, pesanteur, efforts extérieurs",
            ],
            ["Résultats", "Capteurs, Courbes, Animation", "Grandeurs calculées et restitution"],
        ],
        [40, 56, 70],
    )
    g.p(
        "Un élément incomplet est signalé par le symbole **⚠** et un texte orange : liaison non "
        "validée, graphe hyperstatique (h > 0), trop de lois d'entrée. La marque remonte à la "
        "rubrique et à la branche qui le contiennent, ce qui indique à l'élève ce qui reste à "
        "faire."
    )
    g.h2("Les menus")
    g.tableau(
        ["Menu", "Commandes"],
        [
            [
                "Fichier",
                "Nouvelle étude, Ouvrir une étude…, Enregistrer, Enregistrer sous…, Exporter les résultats en CSV…, Quitter",
            ],
            ["SolidCraft", "Connecter, Importer la maquette"],
            ["Simulation", "Simuler (dynamique), Simuler (cinématique), Annuler le calcul"],
            ["Enseignant", "Validation en bloc, Hyperstatique autorisé (chapitre 8)"],
        ],
        [30, 136],
    )
    g.h2("Le journal")
    g.p(
        "L'onglet **Journal** conserve l'historique des messages de la session : import, "
        "validations, refus du garde-fou, fin de calcul, erreurs. C'est le premier endroit à "
        "consulter quand une commande ne produit pas l'effet attendu."
    )
    g.accent()

    # ------------------------------------------------------------------------------ 5
    g.h1("Conduire une étude complète", "05 — Démarche")
    g.p(
        "Ce chapitre suit l'ordre d'une étude réelle, de la maquette SolidCraft aux torseurs. "
        "Chaque étape correspond à une rubrique de l'arbre."
    )

    g.h2("5.1 Préparer la maquette dans SolidCraft")
    g.puces(
        [
            "Chaque **classe d'équivalence** est une occurrence de premier niveau de "
            "l'assemblage. Les sous-occurrences en font partie rigidement.",
            "Le **bâti** est l'occurrence fixée (« grounded ») dans SolidCraft.",
            "Les **liaisons SolidCraft** (joints) relient deux occurrences de premier niveau. "
            "Une liaison sur une sous-occurrence est signalée et ignorée.",
            "Les **couplages** se définissent dans SolidCraft : engrenage (`joint.gear`), "
            "pignon-crémaillère (`joint.rack_pinion`), vis-écrou (`joint.screw`). Meca3D les "
            "importe avec leur rapport.",
            "Les masses viennent des matériaux SolidCraft : affecter un matériau à chaque pièce.",
            "Placer la maquette dans une **configuration générale**, hors point mort : le calcul "
            "de m et h se fait dans la pose d'import.",
        ]
    )

    g.h2("5.2 Connexion et import")
    g.p(
        "**SolidCraft → Connecter** établit la liaison avec SolidCraft ; la barre d'état affiche "
        "« SolidCraft : connecté (port N) ». **SolidCraft → Importer la maquette** lit ensuite "
        "les occurrences, les liaisons, les couplages, les masses, les centres de gravité et les "
        "matrices d'inertie. L'import se fait dans la pose courante de la maquette, qui devient "
        "la configuration de référence de l'étude."
    )
    g.p(
        "Si la maquette change après l'import (pièce ajoutée, liaison modifiée, masse changée), "
        "un bandeau « Maquette modifiée » propose de **Réimporter**. Les choix de l'élève sont "
        "conservés pour les liaisons de même nom."
    )

    g.h2("5.3 Choisir et valider les liaisons")
    g.p(
        "Pour chaque liaison SolidCraft, Meca3D propose une liaison normalisée. L'élève la "
        "confirme ou choisit l'une des onze liaisons du catalogue, puis clique sur **Valider "
        "cette liaison**. Aucune simulation n'est possible tant qu'une liaison reste à valider."
    )
    g.tableau(
        ["Liaison SolidCraft", "Proposition Meca3D"],
        [
            ["revolute", "pivot"],
            ["slider", "glissière"],
            ["cylindrical", "pivot glissant"],
            ["cylindrical couplée par une vis", "hélicoïdale"],
            ["planar", "appui plan"],
            ["ball", "rotule"],
            ["rigid", "encastrement"],
            ["pin_slot", "aucune : choix imposé à l'élève"],
        ],
        [70, 96],
        "Propositions faites à l'import",
    )
    g.tableau(
        ["Liaison normalisée", "ddl bloqués [Tx Ty Tz Rx Ry Rz]", "Axe utilisé"],
        [
            ["encastrement", "[1 1 1 1 1 1]", "—"],
            ["pivot", "[1 1 1 0 1 1]", "axe de rotation"],
            ["glissière", "[0 1 1 1 1 1]", "direction de translation"],
            ["pivot glissant", "[0 1 1 0 1 1]", "axe"],
            ["rotule à doigt", "[1 1 1 0 0 1]", "axe + seconde direction"],
            ["rotule", "[1 1 1 0 0 0]", "—"],
            ["appui plan", "[0 0 1 1 1 0]", "normale au plan"],
            ["linéaire annulaire", "[0 1 1 0 0 0]", "axe"],
            ["linéaire rectiligne", "[0 0 1 0 1 0]", "normale + seconde direction"],
            ["ponctuelle", "[0 0 1 0 0 0]", "normale"],
            ["hélicoïdale", "[0 1 1 0 1 1] et Tx = p·Rx/2π", "axe, pas p"],
        ],
        [44, 66, 56],
        "Catalogue des onze liaisons, en base de liaison (x = axe, z = normale)",
    )
    g.figure(CAPTURES / "g02_liaison.png", "Propriétés d'une liaison : type, statut, axe et centre")
    g.p(
        "Le panneau de propriétés permet aussi de **réorienter l'axe** de la liaison (axe x, y ou "
        "z du repère monde, ou vecteur saisi) et de **déplacer son centre**. C'est indispensable "
        "quand l'axe issu de SolidCraft ne convient pas : par exemple une linéaire annulaire "
        "tirée d'une rotule. La rotule à doigt et la linéaire rectiligne demandent une "
        "**seconde direction**, projetée orthogonalement à l'axe ; une direction presque "
        "colinéaire à l'axe (moins de 1°) est refusée."
    )

    g.h2("5.4 Graphe, mobilité et hyperstatisme")
    g.p(
        "La rubrique **Graphe** affiche le graphe des liaisons : les solides sont les nœuds, les "
        "liaisons les arêtes, le bâti est encadré. Un cycle choisi dans la liste est mis en "
        "évidence en rouge. Les couplages sont tracés en tireté cyan, avec leur rapport."
    )
    g.p(
        "Meca3D calcule μ = L − S + 1 par composante connexe, puis applique la méthode "
        "cinématique : Ic inconnues cinématiques, Ec = 6μ équations (plus une par couplage), "
        "rc rang du système. On en déduit **m = Ic − rc** et **h = Ec − rc**. Le calcul est "
        "refait à chaque changement de liaison."
    )
    g.figure(
        CAPTURES / "g03_hyperstatique.png",
        "Bielle-manivelle avec une rotule en B : h = 1, graphe signalé",
    )
    g.p(
        "Avec les propositions initiales (pivot, pivot, rotule, glissière), la bielle-manivelle "
        "est hyperstatique d'ordre 1. En remplaçant la rotule en B par une **linéaire "
        "annulaire d'axe z**, l'élève obtient un modèle isostatique (h = 0). C'est l'exercice de "
        "modélisation central de Meca3D."
    )
    g.encadre(
        "Garde-fou avant simulation",
        "Si h > 0, la simulation est refusée : le message donne h et les liaisons du cycle "
        "concerné. En mode cinématique, il faut exactement m lois d'entrée. En mode dynamique, "
        "moins de m lois est admis : les mobilités non pilotées évoluent librement (pendule).",
    )
    g.encadre(
        "Configuration singulière",
        "Au point mort d'un mécanisme, le rang chute et m et h sont faux pour cette pose. Meca3D "
        "évalue le rang dans des configurations voisines et prévient l'élève si le résultat "
        "varie. Importer de préférence la maquette hors point mort.",
        CUIVRE,
    )

    g.h2("5.5 Lois d'entrée (pilotes)")
    g.p(
        "Une loi d'entrée impose le mouvement d'une liaison pilote : angle d'un pivot ou d'une "
        "hélicoïdale, position d'une glissière, angle ou position d'un pivot glissant. Elle se "
        "saisit dans le panneau de la liaison, groupe **Loi d'entrée (pilote)**, comme une "
        "expression de t. Les vitesses et accélérations sont obtenues par dérivation symbolique "
        "exacte. L'aperçu affiche valeur, vitesse et accélération à l'instant choisi."
    )
    g.tableau(
        ["Élément", "Syntaxe admise"],
        [
            ["Opérations", "`+  -  *  /  **`"],
            ["Fonctions", "`sin cos tan exp log sqrt abs min max`"],
            ["Constantes", "`pi`, t (temps, s)"],
            ["Loi par morceaux", "`Piecewise((expr1, condition1), (expr2, True))`"],
        ],
        [40, 126],
        "Grammaire des lois d'entrée",
    )
    g.p("Exemples :")
    g.code(
        [
            "20*t                       rotation à 20 rad/s",
            "0.05*sin(2*pi*t)           oscillation de 50 mm à 1 Hz",
            "Piecewise((10*(t - 0.2/pi*sin(pi*t/0.2)), t < 0.2), (2 + 20*(t - 0.2), True))",
            "                           montée en vitesse sur 0,2 s, puis 20 rad/s",
        ]
    )
    g.p(
        "Tout autre nom est refusé, ainsi que l'expression vide. Une montée en vitesse "
        "progressive évite les accélérations infinies au démarrage."
    )
    g.figure(CAPTURES / "g04_loi_entree.png", "Saisie d'une loi d'entrée, avec aperçu à t = 0,5 s")

    g.h2("5.6 Pesanteur, actions et ressorts")
    g.puces(
        [
            "**Pesanteur** : vecteur g en repère monde, par défaut (0 ; 0 ; −9,81) m/s² (z vers "
            "le haut, comme SolidCraft). La case « Pesanteur active » la désactive.",
            "**Actions** : force (N) ou couple (N·m) appliqué en un point d'un solide, exprimé en "
            "repère monde ou en repère du solide. Chaque composante est une constante ou une "
            "fonction de t (même grammaire que les lois d'entrée). Le point d'application se "
            "saisit en coordonnées pièce, en mm.",
            "**Ressorts** : ressort-amortisseur linéaire entre un point de deux solides, avec "
            "raideur k (N/m), longueur libre l₀ (mm) et amortissement c (N·s/m).",
        ]
    )
    g.p(
        "Pour créer une action ou un ressort, sélectionner la rubrique, remplir le panneau et "
        "cliquer sur **Ajouter**. Pour modifier, sélectionner l'élément puis **Appliquer les "
        "modifications** ; **Supprimer** le retire."
    )

    g.h2("5.7 Capteurs")
    g.p("Les capteurs définissent les grandeurs à calculer et à tracer :")
    g.tableau(
        ["Capteur", "Grandeurs", "Options"],
        [
            [
                "Capteur de point",
                "position, vitesse et accélération d'un point d'un solide",
                "solide de référence quelconque ; base du solide de référence ou base monde",
            ],
            [
                "Capteur de liaison",
                "torseur statique (dynamique seulement) et torseur cinématique d'une liaison",
                "sens i → j ou j → i ; base de liaison ou base monde",
            ],
        ],
        [36, 66, 64],
    )
    g.figure(
        CAPTURES / "g05_capteur.png",
        "Capteur de point B du piston, relativement au bâti, en base monde",
    )

    g.h2("5.8 Simuler")
    g.tableau(
        ["Mode", "Principe", "Résultats"],
        [
            [
                "Dynamique",
                "Équations du mouvement intégrées par Exudyn (generalized-α, ρ∞ = 0,9) à pas fixe",
                "mouvement, torseurs de liaison, efforts des couplages",
            ],
            [
                "Cinématique",
                "Mouvement imposé par m lois d'entrée, équations de liaison résolues et dérivées exactement",
                "positions, vitesses, accélérations ; aucun effort",
            ],
        ],
        [28, 80, 58],
    )
    g.p(
        "Le calcul tourne dans un processus séparé : l'interface reste utilisable, la barre "
        "d'état montre la progression, et **Annuler le calcul** l'arrête en moins d'une seconde, "
        "sans résultat partiel. En cas d'échec du solveur (configuration singulière, mécanisme "
        "bloqué), le message indique l'instant de l'échec et rappelle les causes possibles."
    )
    g.p(
        "Les **réglages de simulation** (durée, nombre de pas, période d'écriture des capteurs) "
        "sont enregistrés dans l'étude. Une étude importée reçoit 1 s, 1 000 pas et 1 ms. "
        "Dans la version 0.1.0, ils ne se modifient pas encore depuis l'interface (voir "
        "l'annexe B pour les modifier dans le fichier d'étude)."
    )

    g.h2("5.9 Exploiter les résultats")
    g.h3("Courbes")
    g.p(
        "L'onglet **Courbes** trace toute grandeur calculée en fonction du temps ou d'une autre "
        "grandeur : par exemple la trajectoire d'un point en traçant z en fonction de x. Le "
        "minimum, le maximum et la moyenne de l'ordonnée sont affichés."
    )
    g.figure(
        CAPTURES / "g10_pendule_double.png",
        "Pendule double : trajectoire du bout de la seconde barre, z(x)",
    )
    g.h3("Torseurs")
    g.p(
        "L'onglet **Torseurs** donne, à l'instant choisi par le curseur, le torseur statique "
        "{i→j} ou cinématique {V j/i} de chaque liaison, réduit au centre de la liaison, dans la "
        "base de liaison ou en base monde, dans le sens i → j (action du solide a sur le solide "
        "b) ou j → i. Les composantes nulles par construction sont affichées nulles."
    )
    g.figure(
        CAPTURES / "g07_torseurs.png",
        "Torseurs statiques des quatre liaisons à t = 0,5 s, en base monde",
    )
    g.h3("Export CSV")
    g.p(
        "**Fichier → Exporter les résultats en CSV…** écrit toutes les grandeurs : séparateur "
        "point-virgule, virgule décimale, UTF-8 avec BOM, une colonne par grandeur avec son "
        "unité, par exemple `t (s)` ou `C (N·m)`. Le fichier s'ouvre directement dans Excel en "
        "français."
    )
    g.h3("Animation dans SolidCraft")
    g.p(
        "L'onglet **Animation** rejoue le mouvement simulé dans la maquette SolidCraft : lecture, "
        "pause, vitesse de ×0,1 à ×2, curseur temporel. Il faut une simulation terminée, "
        "SolidCraft connecté et la maquette de l'étude ouverte. La pose initiale de la maquette "
        "est restaurée à l'arrêt, en fin de lecture et à la fermeture ; Meca3D n'enregistre "
        "jamais la maquette."
    )
    g.p(
        "**Tracer la trajectoire** dessine dans SolidCraft, sous forme d'esquisse 3D, la "
        "trajectoire d'un capteur de point. C'est la seule écriture de Meca3D dans la maquette ; "
        "**Effacer les trajectoires** la retire."
    )
    g.figure(
        CAPTURES_LOTS / "lot10_animation.png",
        "Animation de la bielle-manivelle dans SolidCraft, curseur à t = 0,5 s",
    )

    g.h2("5.10 Enregistrer l'étude")
    g.p(
        "**Fichier → Enregistrer** produit un fichier `.meca3d` (JSON, UTF-8). Il contient le "
        "mécanisme importé (solides, masses, inerties, liaisons), les choix de liaison, les "
        "couplages, les lois d'entrée, les actions, les ressorts, les capteurs et les réglages. "
        "Il ne contient pas les résultats, recalculés à chaque simulation de façon identique. "
        "Un fichier corrompu est refusé avec un message, sans étude partielle."
    )
    g.accent()

    # ------------------------------------------------------------------------------ 6
    g.h1("Les exemples fournis", "06 — Cas de référence")
    g.p(
        "Le dossier `exemples` contient cinq études complètes, prêtes à simuler, et les "
        "maquettes SolidCraft correspondantes dans `exemples\\maquettes`. Toutes ont été "
        "construites dans SolidCraft et vérifiées contre des solutions analytiques "
        "indépendantes."
    )
    g.tableau(
        ["Étude", "Mécanisme", "m / h", "Mode", "À observer"],
        [
            [
                "bielle_manivelle",
                "manivelle, bielle, piston",
                "1 / 0",
                "dyn. et ciném.",
                "course du piston, couple moteur, torseurs",
            ],
            [
                "pendule_simple",
                "barre d'acier 400 mm lâchée à 30°",
                "1 / 0",
                "dynamique libre",
                "période 1,054 s, réaction au pivot",
            ],
            [
                "pendule_double",
                "deux barres lâchées à 60°",
                "2 / 0",
                "dynamique libre",
                "mouvement à deux mobilités, trajectoire",
            ],
            [
                "reducteur",
                "engrenage Z1 = 20 / Z2 = 60",
                "1 / 0",
                "dyn. et ciném.",
                "rapport 1/3, couples transmis",
            ],
            [
                "vis_ecrou",
                "vis de pas 5 mm et écrou",
                "1 / 0",
                "dyn. et ciném.",
                "vitesse 7,958 mm/s, effort axial",
            ],
        ],
        [28, 42, 14, 24, 58],
        "Études d'exemple",
    )

    g.h2("Bielle-manivelle")
    g.p(
        "Manivelle de 100 mm, bielle de 300 mm, piston de 0,5 kg, mouvement dans le plan XY avec "
        "une pesanteur g = (0 ; −9,81 ; 0) m/s². La manivelle monte en vitesse sur 0,2 s, puis "
        "tourne à 20 rad/s. La liaison en B est une linéaire annulaire d'axe z, qui rend le "
        "modèle isostatique. Capteurs : point B du piston, point A de la bielle relativement à la "
        "manivelle (vitesse nulle : A est le centre du pivot), torseurs L01, L12 et L30."
    )
    g.p(
        "Exercice suggéré : remettre la liaison en B en rotule, constater h = 1 et le refus de "
        "simulation, puis retrouver la solution isostatique."
    )
    g.h2("Pendule simple")
    g.p(
        "Barre d'acier 400 × 20 × 20 mm (1,248 kg) en pivot d'axe y, lâchée au repos à 30° de la "
        "verticale. Aucune loi d'entrée : le mouvement est libre sous la pesanteur. La période "
        "mesurée est 1,054 s, contre 1,036 s en petites oscillations : l'isochronisme n'est "
        "qu'approché. La vitesse angulaire au point bas vaut 3,139 rad/s."
    )
    g.h2("Pendule double")
    g.p(
        "Deux barres identiques articulées, lâchées à 60° / 60°. Le mécanisme a deux mobilités "
        "(m = 2) et aucune loi d'entrée. Le mouvement est très sensible aux conditions "
        "initiales ; l'énergie mécanique est conservée à mieux que 10⁻³ près sur 2 s."
    )
    g.h2("Réducteur à engrenage")
    g.figure(
        CAPTURES / "g09_reducteur.png",
        "Réducteur : l'engrenage lu dans SolidCraft apparaît en tireté sur le graphe",
    )
    g.p(
        "Deux roues en pivot sur le bâti, couplées par un engrenage extérieur Z1 = 20 / Z2 = 60 "
        "défini dans SolidCraft. Le graphe n'a aucun cycle (μ = 0), mais le couplage ajoute une "
        "équation : m = 1. La roue d'entrée est pilotée à 30 rad/s ; la roue de sortie tourne à "
        "−10 rad/s contre un couple résistant de 1 N·m. Pesanteur désactivée : les axes sont "
        "verticaux."
    )
    g.h2("Système vis-écrou")
    g.p(
        "Vis en pivot et écrou en glissière, couplés par une liaison vis-écrou de pas 5 mm. La "
        "vis tourne à 10 rad/s : l'écrou avance à p·ω/2π = 7,958 mm/s. L'effort résistant de "
        "100 N sur l'écrou exige un couple moteur de F·p/2π = 0,0796 N·m sans perte."
    )
    g.accent()

    # ------------------------------------------------------------------------------ 7
    g.h1("Conventions", "07 — Repères et unités")
    g.tableau(
        ["Sujet", "Convention"],
        [
            ["Unités internes", "SI : m, kg, s, rad, N, N·m"],
            [
                "Saisie et affichage",
                "longueurs en mm dans les panneaux, angles des lois en rad ; l'unité est toujours affichée",
            ],
            ["Repère monde", "celui de SolidCraft, z vers le haut"],
            ["Base de liaison", "x = axe de la liaison, z = normale (appui plan, ponctuelle)"],
            [
                "Torseur statique {i→j}",
                "action du solide i (solide a de la liaison) sur le solide j (solide b), réduit au centre de la liaison",
            ],
            ["Torseur cinématique {V j/i}", "mouvement du solide j par rapport au solide i"],
            ["Repère d'un solide", "repère du composant SolidCraft"],
            ["Points", "saisis en coordonnées pièce (mm), liés au solide"],
        ],
        [48, 118],
        "Conventions de Meca3D",
    )
    g.accent()

    # ------------------------------------------------------------------------------ 8
    g.h1("Usage enseignant", "08 — Menu Enseignant")
    g.tableau(
        ["Réglage", "Effet", "Par défaut"],
        [
            [
                "Validation en bloc",
                "ajoute le bouton « Valider toutes les propositions » (avec confirmation) ; choix d'ergonomie, sans mot de passe",
                "désactivé",
            ],
            [
                "Hyperstatique autorisé",
                "autorise la simulation **cinématique** pour h > 0 : Meca3D retire h équations redondantes et liste les ddl libérés ; la dynamique reste refusée",
                "désactivé",
            ],
        ],
        [40, 100, 26],
    )
    g.h2("Pistes pédagogiques")
    g.puces(
        [
            "**Modélisation** : partir des propositions automatiques, faire chercher le modèle "
            "isostatique par le changement de liaisons (bielle-manivelle).",
            "**Mobilité** : comparer m et h pour plusieurs choix de liaisons, et justifier par le "
            "calcul à la main (Ic, Ec, rc).",
            "**Statique et dynamique** : comparer les torseurs à vitesse lente et rapide, isoler "
            "la part des effets d'inertie.",
            "**Transmission** : vérifier la conservation de la puissance dans le réducteur et le "
            "système vis-écrou (Ce·ωe = Cs·ωs).",
            "**Corrigés** : préparer l'étude complète, puis distribuer une version aux liaisons "
            "seulement proposées.",
        ]
    )
    g.accent()

    # ------------------------------------------------------------------------------ 9
    g.h1("Ligne de commande", "09 — meca3d-cli")
    g.p(
        "`meca3d-cli.exe` simule une étude enregistrée sans interface ni SolidCraft. Il sert aux "
        "corrigés, aux séries de calcul et aux vérifications."
    )
    g.code(
        [
            "meca3d-cli simuler exemples\\bielle_manivelle.meca3d --csv bielle.csv",
            "meca3d-cli simuler etude.meca3d --csv sortie.csv --mode cinematique",
        ]
    )
    g.tableau(
        ["Code de retour", "Signification"],
        [
            ["0", "simulation réussie, CSV écrit"],
            [
                "2",
                "simulation refusée par le garde-fou (h > 0, nombre de lois incorrect) : message affiché",
            ],
            ["autre", "fichier illisible ou échec du calcul : message affiché"],
        ],
        [36, 130],
    )
    g.p("Depuis les sources, la même commande s'écrit `uv run meca3d simuler …`.")
    g.accent()

    # ------------------------------------------------------------------------------ 10
    g.h1("Messages et dépannage", "10 — Questions fréquentes")
    g.tableau(
        ["Situation", "Cause et remède"],
        [
            [
                "« SolidCraft n'est pas lancé ou n'accepte pas les connexions »",
                "Lancer SolidCraft 0.15.0. Si un port a été imposé, vérifier `--control`.",
            ],
            ["« jeton refusé »", "SolidCraft a été relancé : cliquer de nouveau sur Connecter."],
            [
                "Avertissement de format SolidCraft",
                "Version de SolidCraft différente de 0.15.0 : l'import est tenté mais non garanti.",
            ],
            ["Bouton Simuler grisé", "Une liaison reste à valider (⚠ dans l'arbre)."],
            [
                "« Mobilité indisponible »",
                "Une liaison n'a pas de type, ou une seconde direction manque.",
            ],
            [
                "Simulation refusée, h > 0",
                "Revoir les liaisons du cycle indiqué, ou, en cinématique seulement, cocher Hyperstatique autorisé.",
            ],
            ["Simulation cinématique refusée", "Il faut exactement m lois d'entrée."],
            [
                "Échec du solveur à t = …",
                "Point mort ou blocage : vérifier la pose d'import, les lois d'entrée et la durée.",
            ],
            [
                "Animation lente",
                "La fenêtre SolidCraft est réduite ou masquée : la garder visible.",
            ],
            ["Animation indisponible", "Lancer d'abord une simulation et connecter SolidCraft."],
            ["Bandeau « Maquette modifiée »", "La structure de la maquette a changé : Réimporter."],
            ["Blocage par SmartScreen", "Informations complémentaires → Exécuter quand même."],
        ],
        [64, 102],
    )
    g.accent()

    # ------------------------------------------------------------------------------ 11
    g.h1("Limites de la version 0.1.0", "11 — Périmètre")
    g.h2("Hors périmètre")
    g.puces(
        [
            "saisie d'un mécanisme sans SolidCraft (mode autonome, prévu en V2) ;",
            "contact unilatéral, chocs, décollement ; solides déformables, jeux ;",
            "frottement visqueux, résistance au roulement ;",
            "trains épicycloïdaux (couplage porté par un porte-satellites mobile) ;",
            "autres CAO que SolidCraft 0.15.0 ; autres systèmes que Windows ;",
            "post-traitement avancé (FFT) : seuls minimum, maximum et moyenne sont fournis.",
        ]
    )
    g.h2("Réglages non encore exposés dans l'interface")
    g.p(
        "Le moteur de calcul et le fichier d'étude les prennent en charge, mais l'interface de la "
        "version 0.1.0 ne permet pas encore de les saisir :"
    )
    g.puces(
        [
            "durée, nombre de pas et période d'écriture de la simulation ;",
            "coefficient de frottement de Coulomb d'une liaison et rayon d'axe d'un pivot frottant ;",
            "rendement d'un couplage, et saisie d'un couplage absent de la maquette.",
        ]
    )
    g.p("L'annexe B indique comment les renseigner dans le fichier `.meca3d`.")
    g.h2("Distribution")
    g.puces(
        [
            "exécutable non signé : avertissement SmartScreen à la première ouverture ;",
            "premier calcul plus lent (environ 5 s de démarrage du processus de calcul), les "
            "suivants démarrent en 2 s environ.",
        ]
    )
    g.accent()

    # ------------------------------------------------------------------------------ annexes
    g.h1("Annexe A — Glossaire", "Annexes")
    g.tableau(
        ["Terme", "Définition"],
        [
            [
                "Classe d'équivalence",
                "ensemble de pièces sans mouvement relatif ; une occurrence de premier niveau SolidCraft",
            ],
            ["Liaison normalisée", "modèle de liaison parfaite de la norme NF EN ISO 3952"],
            ["Liaison pilote", "liaison dont le mouvement est imposé par une loi d'entrée"],
            ["μ (nombre cyclomatique)", "nombre de cycles indépendants du graphe : μ = L − S + 1"],
            [
                "Ic, Ec, rc",
                "inconnues cinématiques, équations cinématiques (6μ + couplages), rang du système",
            ],
            ["m (mobilité)", "nombre de mouvements indépendants : m = Ic − rc"],
            ["h (hyperstatisme)", "nombre d'équations redondantes : h = Ec − rc"],
            [
                "Torseur statique",
                "actions mécaniques transmises par une liaison : résultante et moment",
            ],
            ["Torseur cinématique", "mouvement relatif : vitesse angulaire et vitesse d'un point"],
            [
                "Couplage",
                "relation imposée entre deux liaisons : engrenage, pignon-crémaillère, vis-écrou",
            ],
            [
                "Generalized-α",
                "intégrateur temporel implicite d'Exudyn ; ρ∞ règle l'amortissement numérique",
            ],
        ],
        [44, 122],
    )

    g.h1("Annexe B — Le fichier d'étude", "Annexes")
    g.p(
        "Le fichier `.meca3d` est un texte JSON lisible, en UTF-8. Un enseignant peut le modifier "
        "avec un éditeur de texte (Bloc-notes, VS Code) pour renseigner les réglages que "
        "l'interface 0.1.0 n'expose pas. Enregistrer une copie avant toute modification : un "
        "fichier invalide est refusé à l'ouverture."
    )
    g.h3("Réglages de simulation")
    g.code(['"reglages": {"duree": 3.0, "nb_pas": 3000, "periode_ecriture": 0.005}'])
    g.p(
        "Durée en s, nombre de pas d'intégration, période d'écriture des capteurs en s. Un pas "
        "de 1 ms convient à la plupart des mécanismes de TP."
    )
    g.h3("Frottement de Coulomb d'une liaison")
    g.code(
        [
            '{"nom": "L30_glissiere", ..., "frottement": 0.2}',
            '{"nom": "L01_pivot", ..., "frottement": 0.1, "rayon_axe": 0.005}',
        ]
    )
    g.p(
        "Coefficient f sans dimension ; rayon d'axe en m pour un pivot. La loi est régularisée "
        "près de la vitesse nulle (v_ε = 10⁻⁴ m/s, ω_ε = 10⁻² rad/s)."
    )
    g.h3("Rendement d'un couplage")
    g.code(['{"genre": "engrenage", ..., "k": 0.3333333333333333, "rendement": 0.95}'])
    g.p(
        "Rendement η entre 0 (exclu) et 1. Le sens de la perte suit le signe de la puissance "
        "transmise à chaque instant."
    )
    g.h3("Structure générale")
    g.tableau(
        ["Clé", "Contenu"],
        [
            ["`version_etude`", "version du format (2)"],
            ["`maquette`, `revision`", "maquette SolidCraft d'origine"],
            ["`solides`", "nom, fixe, masse (kg), centre de gravité (m), inertie (kg·m²), pose"],
            ["`liaisons`", "liaison SolidCraft, type choisi, axe, statut, frottement"],
            ["`couplages`", "genre, liaisons couplées, rapport, rendement"],
            ["`pilotes`", "liaison et loi d'entrée"],
            ["`pesanteur`, `actions`, `ressorts`", "données d'entrée"],
            ["`capteurs`", "capteurs de point et de liaison"],
            ["`reglages`", "durée, nombre de pas, période d'écriture"],
        ],
        [56, 110],
    )

    g.h1("Annexe C — Références et licence", "Annexes")
    g.puces(
        [
            "Code source, versions et suivi : github.com/nbonzani/Meca3D.",
            "Spécification détaillée : `docs/SPEC.md` du dépôt (exigences EX-01 à EX-64).",
            "Exudyn, moteur multicorps : J. Gerstmayr, Université d'Innsbruck.",
            "Decade, logiciel de T. Paviot : source d'inspiration de la démarche pédagogique.",
            "SolidCraft : modeleur CAO piloté par son canal de contrôle.",
        ]
    )
    g.p(
        "Meca3D est un logiciel libre distribué sous licence **GNU GPL version 3 ou "
        "ultérieure**. Les licences des bibliothèques embarquées (Exudyn, NumPy, Matplotlib, "
        "PyQt6, NetworkX, SymPy) sont reproduites dans le dossier `licences` de l'exécutable."
    )
    g.accent()


def main() -> None:
    g = Guide()
    contenu(g)
    g.enregistrer()
    print(f"Guide écrit : {SORTIE}")


if __name__ == "__main__":
    main()
